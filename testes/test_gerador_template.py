import tempfile
import unittest
import zipfile
from pathlib import Path
from openpyxl import Workbook, load_workbook
from src.gerador_template import gerar_rascunho


class TesteGeradorTemplate(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.origem = Path(self.temp.name) / "template.xlsm"
        self.saida = Path(self.temp.name) / "saida.xlsm"
        wb = Workbook()
        ws = wb.active
        ws.title = "Modelo"
        ws["A4"] = "SKU"
        ws["B4"] = "Nome do item"
        ws["A6"] = "EXEMPLO_NAO_COPIAR"
        ws["A6"].number_format = "@"
        wb.create_sheet("Valores válidos")["A1"] = "preservado"
        wb.save(self.origem)

    def test_preserva_formato_e_abas(self):
        gerar_rascunho(self.origem, self.saida, [{"SKU": "1234-39", "Nome do item": "Tênis & teste"}])
        wb = load_workbook(self.saida)
        ws = wb["Modelo"]
        self.assertEqual(ws["A7"].value, "1234-39")
        self.assertEqual(ws["B7"].value, "Tênis & teste")
        self.assertEqual(ws["A6"].value, "EXEMPLO_NAO_COPIAR")
        self.assertEqual(ws["A7"].style_id, ws["A6"].style_id)
        self.assertEqual(wb["Valores válidos"]["A1"].value, "preservado")
        with zipfile.ZipFile(self.origem) as before, zipfile.ZipFile(self.saida) as after:
            for part in before.namelist():
                if part != "xl/worksheets/sheet1.xml":
                    self.assertEqual(before.read(part), after.read(part))

    def test_rejeita_campo_inexistente(self):
        with self.assertRaises(ValueError):
            gerar_rascunho(self.origem, self.saida, [{"Campo inexistente": "X"}])

    def test_nao_sobrescreve_template(self):
        with self.assertRaises(ValueError):
            gerar_rascunho(self.origem, self.origem, [{"SKU": "1234-39"}])

    def test_rejeita_template_preenchido(self):
        wb = load_workbook(self.origem)
        wb["Modelo"]["A7"] = "existente"
        wb.save(self.origem)
        with self.assertRaises(ValueError):
            gerar_rascunho(self.origem, self.saida, [{"SKU": "1234-39"}])


if __name__ == "__main__":
    unittest.main()
