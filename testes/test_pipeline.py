import csv
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from src.pipeline import processar


class TestePipeline(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.csv = self.base / "produtos.csv"
        self.template = self.base / "modelo.xlsm"
        self.saida = self.base / "rascunho.xlsm"
        self.relatorio = self.base / "auditoria.csv"
        wb = Workbook()
        ws = wb.active
        ws.title = "Modelo"
        for i, nome in enumerate(("SKU", "Nome da marca", "Nome do item", "Cor", "Tamanho do calçado"), 1):
            ws.cell(4, i, nome)
        ws["A6"] = "EXEMPLO_NAO_COPIAR"
        wb.save(self.template)

    def gravar(self, linhas):
        with self.csv.open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, delimiter=";", fieldnames=["sku", "categoria", "marca", "nome", "cor", "tamanho"])
            w.writeheader()
            w.writerows(linhas)

    def test_fluxo_completo_e_retencao(self):
        self.gravar([
            {"sku": "1234-39", "categoria": "Calçados", "marca": "Marca Fictícia",
             "nome": "Tênis Exemplo", "cor": "Azul", "tamanho": "39"},
            {"sku": "5678-01", "categoria": "Vestuário", "marca": "Marca Fictícia",
             "nome": "Camiseta Exemplo", "cor": "Branco", "tamanho": "M"},
        ])
        resumo = processar(self.csv, self.template, self.saida, self.relatorio)
        self.assertEqual((resumo["solicitados"], resumo["rascunhos"], resumo["retidos"]), (2, 1, 1))
        wb = load_workbook(self.saida)
        self.assertEqual(wb["Modelo"]["A7"].value, "1234-39")
        self.assertEqual(wb["Modelo"]["A6"].value, "EXEMPLO_NAO_COPIAR")
        self.assertEqual(wb["Modelo"]["D7"].value, "Azul")
        self.assertIsNone(wb["Modelo"]["A8"].value)
        self.assertIn("VARIACAO_REQUER_CONFERENCIA", self.relatorio.read_text(encoding="utf-8-sig"))

    def test_rejeita_duplicado(self):
        item = {"sku": "1234-39", "categoria": "Calçados", "marca": "Marca Fictícia",
                "nome": "Tênis Exemplo", "cor": "Azul", "tamanho": "39"}
        self.gravar([item, item])
        resumo = processar(self.csv, self.template, self.saida, self.relatorio)
        self.assertEqual(resumo["retidos"], 1)
        self.assertIn("SKU_DUPLICADO", self.relatorio.read_text(encoding="utf-8-sig"))

    def test_sem_itens_elegiveis(self):
        self.gravar([{"sku": "1234-39", "categoria": "Vestuário", "marca": "X",
                      "nome": "Camiseta", "cor": "Azul", "tamanho": "M"}])
        resumo = processar(self.csv, self.template, self.saida, self.relatorio)
        self.assertIsNone(resumo["saida"])
        self.assertFalse(self.saida.exists())
        self.assertTrue(self.relatorio.exists())


if __name__ == "__main__":
    unittest.main()
